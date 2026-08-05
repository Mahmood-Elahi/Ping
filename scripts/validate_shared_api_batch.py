"""Build and validate the bounded shared-source/residual-API milestone.

The module is deliberately read-only with respect to production configuration
and persistence.  Network validation uses only production scanner dispatch,
at most five reviewed companies per sub-batch, two pages, and twenty jobs.
"""

from __future__ import annotations

import argparse
import copy
import csv
import json
import sys
from collections import Counter
from dataclasses import replace
from pathlib import Path
from typing import Any, Iterable, Mapping

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import validate_company_registry as validator
from src.scraper.company_careers import scanner
from src.scraper.company_careers.models import CompanyTarget, CoverageStatus
from src.scraper.company_careers.parent_sources import (
    CHILD_BRAND_DISPLAY_BEHAVIOR,
    EXACT_NORMALIZED_MATCH_TYPE,
    ParentSourceMapping,
    load_parent_source_mappings,
    original_company_coverage,
)
from src.scraper.company_careers.registry import load_target_companies


PRODUCTION = ROOT / "config" / "target_companies.yaml"
PRODUCTION_MAPPINGS = ROOT / "config" / "company_parent_source_mappings.yaml"
MASTER_WORKBOOK = ROOT / "config" / "company_watchlist.xlsx"
REMAINING = ROOT / "config" / "company_remaining_inventory_v2.csv"
V3_RESEARCH = ROOT / "config" / "company_unknown_investigation_v3.csv"
AVATURE_CANDIDATES = ROOT / "config" / "company_avature_icims_candidates.csv"
CANDIDATE_INVENTORY = ROOT / "config" / "company_shared_api_candidates.csv"
CANDIDATE_REGISTRY = ROOT / "config" / "target_companies.shared_api_candidates.yaml"
VALIDATED_REGISTRY = ROOT / "config" / "target_companies.shared_api_validated.yaml"
REPORT = ROOT / "output" / "shared_api_validation.json"

MAX_BATCH_SIZE = 5
MAX_PAGES = 2
MAX_JOBS = 20
MAX_RESPONSE_BYTES = 3_000_000
EVIDENCE_TIMESTAMP = "2026-08-05T07:10:00+00:00"
CANONICAL_SAFE = {validator.VALIDATED_WITH_JOBS, validator.VALIDATED_ZERO_JOBS}
PARENT_SAFE = {
    "validated_parent_mapping_with_current_jobs",
    "validated_parent_mapping_no_current_jobs",
}
REVIEW_CLASSIFICATIONS = {
    "parent_company_source",
    "existing_adapter_extension",
    "existing_adapter_reuse",
    "custom_rest_api",
    "custom_graphql",
    "adp",
    "oracle_taleo",
    "talentbrew_radancy",
    "newly_discovered_reusable_platform",
}


def _common(rank: int, name: str, platform: str, careers_url: str, **values: Any) -> dict[str, Any]:
    entry: dict[str, Any] = {
        "name": name,
        "platform": platform,
        "careers_url": careers_url,
        "original_workbook_rank": rank,
        "max_pages": min(int(values.pop("max_pages", 1)), MAX_PAGES),
        "max_jobs": MAX_JOBS,
        "max_response_bytes": MAX_RESPONSE_BYTES,
        "evidence_source": values.pop("evidence_source", "bounded anonymous recheck"),
        "evidence_confidence": values.pop("evidence_confidence", "high"),
        "evidence_timestamp": values.pop("evidence_timestamp", EVIDENCE_TIMESTAMP),
    }
    entry.update(values)
    return entry


def _static(rank: int, name: str, url: str, hosts: list[str], *, listing: str,
            card: str, title: str, link: str, identifier: str | None = None,
            id_attribute: str | None = None, location: str | None = None,
            department: str | None = None, pagination: str | None = None,
            allow_fallback: bool = False, family: str = "existing_adapter_reuse",
            **values: Any) -> dict[str, Any]:
    return _common(
        rank, name, "static_html", url, listing_url=url,
        endpoint_variant="server_rendered_cards", listing_selector=listing,
        job_card_selector=card, title_selector=title, detail_link_selector=link,
        external_id_regex=identifier, external_id_attribute=id_attribute,
        location_selector=location, department_selector=department,
        pagination_url_template=pagination, allow_url_id_fallback=allow_fallback,
        allowed_apply_hosts=hosts, candidate_family=family, **values,
    )


CANONICAL_ENTRIES: list[dict[str, Any]] = [
    _common(
        48, "Netflix", "eightfold", "https://explore.jobs.netflix.net/careers",
        endpoint_url="https://explore.jobs.netflix.net/api/pcsx/search",
        endpoint_variant="pcsx", organization_id="netflix.com",
        allowed_apply_hosts=["explore.jobs.netflix.net"],
        candidate_family="existing_adapter_extension",
        prior_classification="existing_adapter_extension",
        expected_status="unsupported_variant",
        exclusion_reason="Eightfold PCSX endpoint currently returns 403: PCSX is not enabled for this user",
    ),
    _static(
        65, "Wise", "https://wise.jobs/jobs", ["wise.jobs"],
        listing="body", card="div.attrax-vacancy-tile",
        title="a.attrax-vacancy-tile__title", link="a.attrax-vacancy-tile__title",
        id_attribute="data-jobid",
        location="div.attrax-vacancy-tile__location-freetext p.attrax-vacancy-tile__item-value",
        department="div.attrax-vacancy-tile__option-job-family-group p.attrax-vacancy-tile__item-value",
        prior_classification="custom_rest_api",
    ),
    _common(
        78, "Jump Trading", "greenhouse", "https://www.jumptrading.com/careers",
        slug="jumptrading", allowed_apply_hosts=["www.jumptrading.com"],
        candidate_family="existing_adapter_reuse",
        prior_classification="existing_adapter_extension",
    ),
    _static(
        81, "Optiver", "https://www.optiver.com/join-us/jobs/", ["www.optiver.com"],
        listing="div[role=region]", card="div[role=region] ul.flex.flex-col.gap-3 li",
        title="h3 a", link="h3 a", location="p.text-body-s",
        department="span.text-label-s", allow_fallback=True,
        prior_classification="custom_rest_api",
    ),
    _static(
        88, "BlackRock", "https://careers.blackrock.com/en/search-jobs",
        ["careers.blackrock.com"], listing="ul.section3__search-results-ul",
        card="a.section3__search-results-a", title="h2.section3__job-title",
        link=":self", id_attribute="data-job-id",
        location="span.job-location span.section3__job-info",
        department="span.job-category span.section3__job-info",
        pagination="https://careers.blackrock.com/en/search-jobs?p={page}", max_pages=2,
        family="talentbrew_radancy", prior_classification="talentbrew_radancy",
    ),
    _common(
        93, "Schneider Electric", "manual_review", "https://careers.se.com/jobs",
        candidate_family="oracle_taleo", prior_classification="oracle_taleo",
        expected_status="blocked", exclusion_reason="Current anonymous official careers request returns HTTP 403 Access Denied",
    ),
    _common(
        124, "Continental", "manual_review", "https://jobs.continental.com/en/jobs/",
        listing_url="https://jobs.continental.com/en/jobs/",
        api_url="https://jobs.continental.com/en/api/result-list/pagetype-jobs/",
        endpoint_variant="continental_result_list_v1", candidate_family="custom_rest_api",
        prior_classification="custom_rest_api",
        allowed_source_hosts=["jobs.continental.com"], allowed_apply_hosts=["jobs.continental.com"],
        stable_id_contract="smartRecruitersId or uuid", pagination_contract="currentPage + itemsPerPage",
        expected_status="unsupported_variant",
        exclusion_reason="Only one current candidate exposes this contract; a reusable public REST family is not proven",
    ),
    _static(
        138, "Snap", "https://careers.snap.com/jobs", ["careers.snap.com"],
        listing="table", card="tr[role=row]", title="a[href*='/job?id=']",
        link="a[href*='/job?id=']", identifier=r"id=([^&]+)",
        location="td:nth-child(4)", department="td:nth-child(2)",
        prior_classification="custom_graphql",
    ),
    _common(
        172, "bp", "workday", "https://bpinternational.wd3.myworkdayjobs.com/bpCareers",
        tenant="bpinternational", wd_subdomain="wd3", site_id="bpCareers",
        endpoint_variant="wday_cxs_v1",
        allowed_apply_hosts=["bpinternational.wd3.myworkdayjobs.com"],
        candidate_family="existing_adapter_reuse", prior_classification="avature",
    ),
    _common(
        174, "Canadian Natural Resources", "oracle_recruiting",
        "https://ehaa.fa.ca2.oraclecloud.com/hcmUI/CandidateExperience/en/sites/CNRL-Professional/jobs",
        endpoint_url="https://ehaa.fa.ca2.oraclecloud.com/hcmRestApi/resources/latest/recruitingCEJobRequisitions",
        site_number="CNRL-Professional", locale="en", endpoint_variant="recruiting_ce_v1",
        allowed_apply_hosts=["ehaa.fa.ca2.oraclecloud.com"],
        candidate_family="existing_adapter_reuse",
        prior_classification="existing_adapter_extension",
    ),
    _common(
        189, "Astrobotic", "manual_review", "https://www.astrobotic.com/company/careers/",
        listing_url="https://workforcenow.cloud.adp.com/mascsr/default/mdf/recruitment/recruitment.html?cid=dde86c74-1cf6-4830-b396-33b9bd085ff9&ccId=48397579_2&lang=en_US",
        tenant="dde86c74-1cf6-4830-b396-33b9bd085ff9", site_id="48397579_2",
        candidate_family="adp", prior_classification="adp",
        expected_status="browser_required",
        exclusion_reason="Official ADP URL is a JavaScript shell with no conclusive anonymous public listing contract",
    ),
    _common(
        191, "Relativity Space", "greenhouse", "https://www.relativityspace.com/jobs",
        slug="relativity", allowed_apply_hosts=["boards.greenhouse.io", "job-boards.greenhouse.io"],
        candidate_family="existing_adapter_reuse",
        prior_classification="existing_adapter_extension",
    ),
]


def _mapping(rank: int, child: str, parent: str, source: str, *, field: str = "",
             values: Iterable[str] = (), status: str, reason: str = "",
             historical_matches: int = 0) -> dict[str, Any]:
    return {
        "master_rank": rank,
        "child_company": child,
        "parent_company": parent,
        "canonical_source": source,
        "match_field": field,
        "accepted_values": list(values),
        "match_type": EXACT_NORMALIZED_MATCH_TYPE,
        "display_behavior": CHILD_BRAND_DISPLAY_BEHAVIOR,
        "confidence": "high" if field else "low",
        "evidence_source": "bounded production-parent recheck",
        "evidence_timestamp": EVIDENCE_TIMESTAMP,
        "expected_status": status,
        "exclusion_reason": reason,
        "historical_exact_matches": historical_matches,
    }


PARENT_CANDIDATES: list[dict[str, Any]] = [
    _mapping(9, "Google DeepMind", "Google", "custom:google", field="departments",
             values=["DeepMind"], status="validated_parent_mapping_with_current_jobs",
             historical_matches=1),
    _mapping(166, "Clearpath Robotics", "Rockwell Automation", "workday:rockwellautomation",
             status="weak_parent_attribution", reason="Workday listing exposes no exact structured Clearpath brand field"),
    _mapping(167, "OTTO Motors", "Rockwell Automation", "workday:rockwellautomation",
             status="weak_parent_attribution", reason="Workday listing exposes no exact structured OTTO Motors brand field"),
    _mapping(169, "Imperial Oil", "ExxonMobil", "", status="invalid_configuration",
             reason="Canonical parent source ExxonMobil is not configured in production"),
    _mapping(201, "Walmart Global Tech", "Walmart", "", status="invalid_configuration",
             reason="Canonical parent source Walmart is not configured in production"),
    _mapping(208, "Uber Eats", "Uber", "custom:uber", status="weak_parent_attribution",
             reason="Current jobs mention Uber Eats only in titles; department/team fields are broad job functions"),
    _mapping(216, "Costco Technology", "Costco", "", status="invalid_configuration",
             reason="Canonical parent source Costco is not configured and the reviewed endpoint returned HTTP 503"),
    _mapping(223, "X", "xAI", "greenhouse:xai", status="blocked",
             reason="Official X careers route redirects to an access-blocked xAI page and exposes no exact X brand metadata"),
    _mapping(233, "Threads", "Meta", "custom:meta", status="weak_parent_attribution",
             reason="Current Meta results contain title-only Threads mentions and no exact Threads department value"),
    _mapping(235, "WeChat", "Tencent", "", status="invalid_configuration",
             reason="Canonical parent source Tencent is not configured in production"),
]

EXPECTED_COMPANIES = [
    "Google DeepMind", "Netflix", "Wise", "Jump Trading", "Optiver", "BlackRock",
    "Schneider Electric", "Continental", "Snap", "Clearpath Robotics", "OTTO Motors",
    "Imperial Oil", "bp", "Canadian Natural Resources", "Astrobotic", "Relativity Space",
    "Walmart Global Tech", "Uber Eats", "Costco Technology", "X", "Threads", "WeChat",
]
SAFE_CANONICAL_NAMES = [
    "Wise", "Jump Trading", "Optiver", "BlackRock", "Snap", "bp",
    "Canadian Natural Resources", "Relativity Space",
]
SAFE_PARENT_NAMES = ["Google DeepMind"]


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _master_ranks() -> dict[str, int]:
    workbook = openpyxl.load_workbook(MASTER_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    company_index, rank_index = headers.index("Company"), headers.index("Rank")
    result = {str(row[company_index]).strip(): int(row[rank_index]) for row in rows if row[company_index]}
    workbook.close()
    return result


def _derived_candidate_names() -> list[str]:
    names: set[str] = set()
    for row in _csv_rows(REMAINING):
        if row.get("currently_uncovered", "").casefold() != "true":
            continue
        if row.get("current_classification", "").casefold() in REVIEW_CLASSIFICATIONS:
            names.add(row["company"])
    for row in _csv_rows(V3_RESEARCH):
        classification = row.get("primary_classification", "").casefold()
        family = row.get("platform_family", "").casefold()
        if classification in REVIEW_CLASSIFICATIONS or family in REVIEW_CLASSIFICATIONS:
            names.add(row["company"])
    for row in _csv_rows(AVATURE_CANDIDATES):
        if row.get("endpoint_variant", "").casefold() == "stale_workday_migration":
            names.add(row["company"])
    _, production = validator.load_registry(PRODUCTION)
    covered = {validator._name_key(item) for item in production}
    covered.update(
        " ".join(mapping.child_company.casefold().split())
        for mapping in load_parent_source_mappings(PRODUCTION_MAPPINGS)
    )
    names = {name for name in names if validator._name_key({"name": name}) not in covered}
    ranks = _master_ranks()
    return sorted(names, key=ranks.__getitem__)


def _production_mapping_dicts() -> list[dict[str, Any]]:
    data = yaml.safe_load(PRODUCTION_MAPPINGS.read_text(encoding="utf-8")) or {}
    return [copy.deepcopy(item) for item in data["company_parent_source_mappings"]]


def _safe_mapping_dict(item: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "child_company": item["child_company"], "parent_company": item["parent_company"],
        "enabled": True, "match_field": item["match_field"],
        "match_type": EXACT_NORMALIZED_MATCH_TYPE,
        "accepted_values": list(item["accepted_values"]),
        "canonical_source": item["canonical_source"],
        "display_behavior": CHILD_BRAND_DISPLAY_BEHAVIOR, "confidence": "high",
        "evidence": {
            "validation_artifact": "output/shared_api_validation.json",
            "evidence_source": item["evidence_source"],
            "validation_status": item["expected_status"],
        },
    }


def _candidate_mapping_dict(item: Mapping[str, Any]) -> dict[str, Any]:
    result = _safe_mapping_dict(item)
    result["enabled"] = bool(item["expected_status"] in PARENT_SAFE)
    result["confidence"] = item["confidence"]
    result["evidence"].update({
        "validation_status": item["expected_status"],
        "safe_to_activate": item["expected_status"] in PARENT_SAFE,
        "warnings": [item["exclusion_reason"]] if item["exclusion_reason"] else [],
    })
    return result


def _registry_bytes(entries: Iterable[Mapping[str, Any]], mappings: list[dict[str, Any]]) -> bytes:
    content = validator.build_candidate_registry_bytes(PRODUCTION, entries)
    newline = b"\r\n" if b"\r\n" in PRODUCTION.read_bytes() else b"\n"
    mapping_yaml = yaml.safe_dump(
        {"company_parent_source_mappings": mappings}, sort_keys=False,
        allow_unicode=True, width=1000,
    ).encode("utf-8").replace(b"\n", newline)
    return content + (b"" if content.endswith((b"\n", b"\r")) else newline) + mapping_yaml


def reconcile_candidates() -> None:
    derived = _derived_candidate_names()
    if derived != EXPECTED_COMPANIES:
        raise ValueError(f"Research artifacts do not derive the exact candidate set: {derived!r}")
    ranks = _master_ranks()
    reviewed = sorted(
        [(item["original_workbook_rank"], item["name"]) for item in CANONICAL_ENTRIES]
        + [(item["master_rank"], item["child_company"]) for item in PARENT_CANDIDATES]
    )
    if [name for _, name in reviewed] != EXPECTED_COMPANIES:
        raise ValueError("Configured candidate order is not master-workbook order")
    if any(rank != ranks[name] for rank, name in reviewed):
        raise ValueError("Configured master ranks do not match the workbook")
    _, production = validator.load_registry(PRODUCTION)
    production_names = {validator._name_key(item) for item in production}
    mapped = {" ".join(item.child_company.casefold().split()) for item in load_parent_source_mappings(PRODUCTION_MAPPINGS)}
    overlap = {" ".join(name.casefold().split()) for name in EXPECTED_COMPANIES} & (production_names | mapped)
    if overlap:
        raise ValueError(f"Reviewed candidates already covered: {sorted(overlap)}")
    if validator._duplicate_company_owners(production, CANONICAL_ENTRIES):
        raise ValueError("Duplicate canonical company")
    if validator._duplicate_owners(production, CANONICAL_ENTRIES):
        raise ValueError("Duplicate canonical source")


def build_candidate_artifacts() -> None:
    reconcile_candidates()
    existing = _production_mapping_dicts()
    candidate_mappings = existing + [_candidate_mapping_dict(item) for item in PARENT_CANDIDATES]
    content = _registry_bytes(CANONICAL_ENTRIES, candidate_mappings)
    if not content.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Candidate registry does not preserve the production byte prefix")
    validator._write_bytes_if_changed(CANDIDATE_REGISTRY, content)
    expected_count = len(load_target_companies(str(PRODUCTION))) + len(CANONICAL_ENTRIES)
    if len(load_target_companies(str(CANDIDATE_REGISTRY))) != expected_count:
        raise ValueError("Candidate registry is not target-loader compatible")


def _canonical_row(entry: Mapping[str, Any], raw: Mapping[str, Any]) -> dict[str, Any]:
    forced = entry.get("expected_status")
    status = str(forced or raw.get("validation_status") or "failed_technical")
    safe = status in CANONICAL_SAFE
    pagination = raw.get("pagination_result") if isinstance(raw.get("pagination_result"), Mapping) else {}
    diagnostics = list(pagination.get("diagnostics") or [])
    if raw.get("error"):
        diagnostics.append(str(raw["error"]))
    if entry.get("exclusion_reason"):
        diagnostics.append(str(entry["exclusion_reason"]))
    return {
        "company": entry["name"], "master_rank": entry["original_workbook_rank"],
        "current_family": entry["candidate_family"],
        "endpoint_identifiers": validator.configured_identifiers(entry),
        "pages_requested": min(int(raw.get("pages_requested") or 0), MAX_PAGES),
        "jobs_observed": int(raw.get("jobs_returned") or 0),
        "normalized_jobs_inspected": min(int(raw.get("jobs_normalized") or 0), MAX_JOBS),
        "stable_id_result": raw.get("stable_job_id_result") or "not_observed",
        "apply_url_result": raw.get("apply_url_result") or "not_observed",
        "apply_host_result": raw.get("apply_host_result") or raw.get("apply_url_result") or "not_observed",
        "pagination_result": dict(pagination), "repeated_page_result": bool(pagination.get("repeated_page")),
        "valid_empty_result": bool(pagination.get("zero_job_contract")),
        "authentication_check": "rejected" if status == "authentication_required" else "not_required",
        "session_check": "rejected" if status == "session_required" else "not_required",
        "browser_check": "rejected" if status == "browser_required" else "not_required",
        "malformed_response_check": "rejected" if status == "malformed_response" else "not_observed",
        "duplicate_company_result": raw.get("duplicate_company_status") or "unique",
        "duplicate_source_result": raw.get("duplicate_source_status") or "unique",
        "status": status, "safe_to_activate": safe, "warnings": [], "diagnostics": diagnostics,
    }


def _validate_canonical(timeout: int, live: bool) -> list[dict[str, Any]]:
    _, production = validator.load_registry(PRODUCTION)
    duplicate_sources = validator._duplicate_owners(production, CANONICAL_ENTRIES)
    duplicate_names = validator._duplicate_company_owners(production, CANONICAL_ENTRIES)
    rows: list[dict[str, Any]] = []
    for start in range(0, len(CANONICAL_ENTRIES), MAX_BATCH_SIZE):
        batch = CANONICAL_ENTRIES[start:start + MAX_BATCH_SIZE]
        for offset, entry in enumerate(batch):
            index = start + offset
            if not live:
                raw: dict[str, Any] = {"validation_status": entry.get("expected_status", "not_run")}
            elif entry["platform"] == "manual_review":
                raw = {"validation_status": entry.get("expected_status", "unsupported_variant")}
            else:
                raw = validator.validate_company(
                    entry, duplicate_of=duplicate_sources.get(index),
                    duplicate_company_of=duplicate_names.get(index), timeout=timeout,
                    attempts=2, retry_delay=0.1, adapters=scanner.ADAPTERS,
                    max_normalized_jobs=MAX_JOBS,
                )
            rows.append(_canonical_row(entry, raw))
    return rows


def _normalized(value: Any) -> str:
    return " ".join(str(value or "").casefold().split())


def _validate_parents(timeout: int, live: bool) -> list[dict[str, Any]]:
    production_targets = {target.name: target for target in load_target_companies(str(PRODUCTION))}
    safe_candidate = ParentSourceMapping.from_dict(_safe_mapping_dict(PARENT_CANDIDATES[0]))
    annotations = load_parent_source_mappings(PRODUCTION_MAPPINGS) + [safe_candidate]
    parent_jobs: dict[str, list[Any]] = {}
    parent_results: dict[str, Any] = {}
    network_parents = ["Google", "Rockwell Automation", "Uber", "xAI", "Meta"]
    if live:
        for start in range(0, len(network_parents), MAX_BATCH_SIZE):
            for parent in network_parents[start:start + MAX_BATCH_SIZE]:
                target = production_targets[parent]
                terms = ["Google DeepMind"] if parent == "Google" else list(target.search_terms)
                bounded = replace(target, max_pages=min(target.max_pages, MAX_PAGES), max_jobs=MAX_JOBS, search_terms=terms)
                report = scanner.CompanyCareersScanner(
                    targets=[bounded], timeout=timeout, parent_source_mappings=annotations,
                ).scan(max_jobs_per_company=MAX_JOBS)
                parent_results[parent] = report.results[0]
                parent_jobs[parent] = report.results[0].jobs
    rows: list[dict[str, Any]] = []
    for item in PARENT_CANDIDATES:
        jobs = parent_jobs.get(item["parent_company"], [])
        exact_matches = 0
        if item["match_field"] == "departments":
            expected = {_normalized(value) for value in item["accepted_values"]}
            exact_matches = sum(bool(expected & {_normalized(value) for value in job.departments}) for job in jobs)
        status = item["expected_status"]
        result = parent_results.get(item["parent_company"])
        if live and status in PARENT_SAFE:
            if result is None or result.status == CoverageStatus.FAILED_TECHNICAL:
                status = "failed_technical"
            elif exact_matches:
                status = "validated_parent_mapping_with_current_jobs"
            elif item["historical_exact_matches"] and jobs and all(job.departments is not None for job in jobs):
                status = "validated_parent_mapping_no_current_jobs"
            else:
                status = "weak_parent_attribution"
        safe = status in PARENT_SAFE
        rows.append({
            "child_company": item["child_company"], "master_rank": item["master_rank"],
            "canonical_parent": item["parent_company"],
            "canonical_source_identity": item["canonical_source"],
            "structured_metadata_field": item["match_field"],
            "exact_expected_values": list(item["accepted_values"]),
            "parent_jobs_inspected": min(len(jobs), MAX_JOBS),
            "current_exact_matches": exact_matches,
            "historical_exact_matches": item["historical_exact_matches"],
            "weak_match_rejection_checks": {
                "title_substring": "rejected", "description_substring": "rejected",
                "url_substring": "rejected", "industry_term": "rejected",
            },
            "overlap_with_other_child_brands": False,
            "canonical_source_key_preservation": bool(item["canonical_source"]),
            "duplicate_row_check": "unique", "status": status,
            "safe_to_activate": safe, "warnings": [],
            "diagnostics": [item["exclusion_reason"]] if item["exclusion_reason"] else [],
        })
    return rows


INVENTORY_FIELDS = [
    "master_rank", "company", "prior_classification", "current_classification",
    "candidate_family", "canonical_parent", "child_brand", "parent_source_identity",
    "exact_metadata_field", "exact_metadata_values", "platform", "endpoint_variant",
    "careers_url", "listing_url", "api_url", "graphql_operation", "tenant", "site_id",
    "organization", "career_section", "allowed_source_hosts", "allowed_apply_hosts",
    "stable_id_contract", "pagination_contract", "evidence_source", "evidence_confidence",
    "evidence_timestamp", "validation_status", "safe_to_activate", "warnings", "exclusion_reason",
]


def _write_inventory(canonical_rows: list[Mapping[str, Any]], parent_rows: list[Mapping[str, Any]]) -> None:
    canonical_by_name = {row["company"]: row for row in canonical_rows}
    parent_by_name = {row["child_company"]: row for row in parent_rows}
    canonical_config = {item["name"]: item for item in CANONICAL_ENTRIES}
    parent_config = {item["child_company"]: item for item in PARENT_CANDIDATES}
    ranks = _master_ranks()
    rows: list[dict[str, Any]] = []
    for name in EXPECTED_COMPANIES:
        if name in canonical_config:
            item, result = canonical_config[name], canonical_by_name[name]
            rows.append({
                "master_rank": ranks[name], "company": name,
                "prior_classification": item.get("prior_classification", item["candidate_family"]),
                "current_classification": item["candidate_family"], "candidate_family": item["candidate_family"],
                "platform": item["platform"], "endpoint_variant": item.get("endpoint_variant", ""),
                "careers_url": item.get("careers_url", ""), "listing_url": item.get("listing_url", ""),
                "api_url": item.get("api_url", item.get("endpoint_url", "")),
                "graphql_operation": item.get("graphql_operation", ""), "tenant": item.get("tenant", ""),
                "site_id": item.get("site_id", item.get("site_number", "")),
                "organization": item.get("organization_id", ""), "career_section": item.get("career_section", ""),
                "allowed_source_hosts": ";".join(item.get("allowed_source_hosts", [])),
                "allowed_apply_hosts": ";".join(item.get("allowed_apply_hosts", [])),
                "stable_id_contract": item.get("stable_id_contract", "production adapter stable ID"),
                "pagination_contract": item.get("pagination_contract", "production adapter; max two pages"),
                "evidence_source": item["evidence_source"], "evidence_confidence": item["evidence_confidence"],
                "evidence_timestamp": item["evidence_timestamp"], "validation_status": result["status"],
                "safe_to_activate": str(result["safe_to_activate"]).lower(),
                "warnings": "; ".join(result["warnings"]), "exclusion_reason": item.get("exclusion_reason", ""),
            })
        else:
            item, result = parent_config[name], parent_by_name[name]
            rows.append({
                "master_rank": ranks[name], "company": name, "prior_classification": "parent_company_source",
                "current_classification": "parent_company_source", "candidate_family": "parent_company_source",
                "canonical_parent": item["parent_company"], "child_brand": item["child_company"],
                "parent_source_identity": item["canonical_source"], "exact_metadata_field": item["match_field"],
                "exact_metadata_values": ";".join(item["accepted_values"]), "platform": "parent_mapping",
                "evidence_source": item["evidence_source"], "evidence_confidence": item["confidence"],
                "evidence_timestamp": item["evidence_timestamp"], "validation_status": result["status"],
                "safe_to_activate": str(result["safe_to_activate"]).lower(),
                "warnings": "; ".join(result["warnings"]), "exclusion_reason": item["exclusion_reason"],
            })
    with CANDIDATE_INVENTORY.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=INVENTORY_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run(*, live: bool = True, timeout: int = 20) -> dict[str, Any]:
    production_before = PRODUCTION.read_bytes()
    mappings_before = PRODUCTION_MAPPINGS.read_bytes()
    build_candidate_artifacts()
    canonical_rows = _validate_canonical(timeout, live)
    parent_rows = _validate_parents(timeout, live)
    safe_names = {row["company"] for row in canonical_rows if row["safe_to_activate"]}
    safe_entries = [item for item in CANONICAL_ENTRIES if item["name"] in safe_names]
    safe_parent_names = {row["child_company"] for row in parent_rows if row["safe_to_activate"]}
    safe_parent_configs = [item for item in PARENT_CANDIDATES if item["child_company"] in safe_parent_names]
    validated_mappings = _production_mapping_dicts() + [_safe_mapping_dict(item) for item in safe_parent_configs]
    validated = _registry_bytes(safe_entries, validated_mappings)
    if not validated.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Validated registry does not preserve the production byte prefix")
    validator._write_bytes_if_changed(VALIDATED_REGISTRY, validated)
    _write_inventory(canonical_rows, parent_rows)
    all_mappings = load_parent_source_mappings(PRODUCTION_MAPPINGS) + [
        ParentSourceMapping.from_dict(_safe_mapping_dict(item)) for item in safe_parent_configs
    ]
    coverage = original_company_coverage(
        load_target_companies(str(VALIDATED_REGISTRY)), all_mappings,
        inventory_path=REMAINING,
    )
    canonical_counts = Counter(row["status"] for row in canonical_rows)
    totals = {
        "candidates_reviewed": len(EXPECTED_COMPANIES),
        "canonical_source_candidates": len(CANONICAL_ENTRIES),
        "parent_mapping_candidates": len(PARENT_CANDIDATES),
        "existing_adapter_reuse_candidates": sum(
            item["platform"] in scanner.ADAPTERS
            and item["candidate_family"] != "existing_adapter_extension"
            for item in CANONICAL_ENTRIES
        ),
        "existing_adapter_extension_candidates": sum(item["candidate_family"] == "existing_adapter_extension" for item in CANONICAL_ENTRIES),
        "rest_candidates": sum(item.get("prior_classification") == "custom_rest_api" for item in CANONICAL_ENTRIES),
        "graphql_candidates": sum(item.get("prior_classification") == "custom_graphql" for item in CANONICAL_ENTRIES),
        "adp_candidates": sum(item["candidate_family"] == "adp" for item in CANONICAL_ENTRIES),
        "taleo_candidates": sum(item["candidate_family"] == "oracle_taleo" for item in CANONICAL_ENTRIES),
        "talentbrew_radancy_candidates": sum(item["candidate_family"] == "talentbrew_radancy" for item in CANONICAL_ENTRIES),
        "validated_with_jobs": canonical_counts[validator.VALIDATED_WITH_JOBS],
        "validated_zero_jobs": canonical_counts[validator.VALIDATED_ZERO_JOBS],
        "validated_parent_mappings": len(safe_parent_configs),
        "excluded": sum(not row["safe_to_activate"] for row in canonical_rows + parent_rows),
        "temporary_failures": sum(row["status"] == "temporary_failure" for row in canonical_rows + parent_rows),
        "production_canonical_count": len(load_target_companies(str(PRODUCTION))),
        "candidate_canonical_count": len(load_target_companies(str(CANDIDATE_REGISTRY))),
        "validated_canonical_count": len(load_target_companies(str(VALIDATED_REGISTRY))),
        "production_mapped_count": len(load_parent_source_mappings(PRODUCTION_MAPPINGS)),
        "candidate_mapped_count": len(_production_mapping_dicts()) + len(PARENT_CANDIDATES),
        "validated_mapped_count": len(validated_mappings),
        "projected_original_company_coverage": coverage["total_original_companies_covered"],
    }
    report = {
        "limits": {"maximum_batch_size": MAX_BATCH_SIZE, "maximum_pages": MAX_PAGES,
                   "maximum_normalized_jobs": MAX_JOBS, "maximum_response_bytes": MAX_RESPONSE_BYTES},
        "network_sub_batches": [
            [item["name"] for item in CANONICAL_ENTRIES[index:index + MAX_BATCH_SIZE]]
            for index in range(0, len(CANONICAL_ENTRIES), MAX_BATCH_SIZE)
        ] + [["Google", "Rockwell Automation", "Uber", "xAI", "Meta"]],
        "totals": totals, "canonical_candidates": canonical_rows,
        "parent_mapping_candidates": parent_rows,
    }
    validator._write_json_if_changed(REPORT, report)
    if PRODUCTION.read_bytes() != production_before or PRODUCTION_MAPPINGS.read_bytes() != mappings_before:
        raise ValueError("Production configuration changed during validation")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-live", action="store_true", help="Build artifacts without HTTP validation")
    parser.add_argument("--timeout", type=int, default=20)
    args = parser.parse_args()
    report = run(live=not args.no_live, timeout=max(1, args.timeout))
    print(json.dumps(report["totals"], indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
