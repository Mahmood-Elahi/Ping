"""Enrich a bounded company-watchlist batch with safe ATS metadata."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

import requests
from openpyxl import load_workbook

try:
    from scripts.company_watchlist_common import (
        ATSDetection,
        ENRICHED_FIELDS,
        canonicalize_url,
        clean_text,
        detect_ats,
        detection_from_registry,
        load_registry,
        normalize_company_name,
        parse_bool,
        read_research_csv,
        registry_by_name,
        registry_source_url,
        validate_detection_fields,
        validate_public_https_url,
        validate_research_row,
        write_csv,
    )
except ModuleNotFoundError:  # Direct execution places scripts/ first on sys.path.
    from company_watchlist_common import (  # type: ignore[no-redef]
        ATSDetection,
        ENRICHED_FIELDS,
        canonicalize_url,
        clean_text,
        detect_ats,
        detection_from_registry,
        load_registry,
        normalize_company_name,
        parse_bool,
        read_research_csv,
        registry_by_name,
        registry_source_url,
        validate_detection_fields,
        validate_public_https_url,
        validate_research_row,
        write_csv,
    )


DEFAULT_INPUT = Path("config/company_watchlist.xlsx")
DEFAULT_OUTPUT = Path("config/company_watchlist_enriched.csv")
DEFAULT_MANUAL_REVIEW = Path("output/company_watchlist_manual_review.csv")
REQUIRED_INPUT_COLUMNS = ["Rank", "Company", "Careers URL", "Priority", "Notes"]


def read_workbook(path: Path | str) -> list[dict[str, str]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(values)
        except StopIteration as exc:
            raise ValueError("Watchlist workbook is empty") from exc
        headers = [clean_text(value) for value in raw_headers]
        missing = [column for column in REQUIRED_INPUT_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"Workbook is missing columns: {', '.join(missing)}")
        rows: list[dict[str, str]] = []
        for excel_row, values_row in enumerate(values, start=2):
            row = {
                header: clean_text(values_row[index] if index < len(values_row) else "")
                for index, header in enumerate(headers)
                if header
            }
            if not any(row.values()):
                continue
            row["_excel_row"] = str(excel_row)
            rows.append(row)
        return rows
    finally:
        workbook.close()


def enrich_rows(
    workbook_rows: list[Mapping[str, Any]],
    registry_entries: list[Mapping[str, Any]],
    *,
    start_row: int = 1,
    limit: Optional[int] = None,
    batch_size: int = 25,
    session: Any = None,
    validate_urls: bool = True,
    research_rows: Optional[list[Mapping[str, Any]]] = None,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    if start_row < 1:
        raise ValueError("start_row must be at least 1")
    if batch_size < 1:
        raise ValueError("batch_size must be at least 1")
    if limit is not None and limit < 1:
        raise ValueError("limit must be at least 1")

    stop = None if limit is None else start_row - 1 + limit
    selected = workbook_rows[start_row - 1 : stop]
    counts = Counter(normalize_company_name(row.get("Company")) for row in workbook_rows)
    existing = registry_by_name(registry_entries)
    reviewed_research = {
        normalize_company_name(row.get("Company")): row
        for row in (research_rows or [])
        if parse_bool(row.get("Reviewed")) and normalize_company_name(row.get("Company"))
    }
    http = session or requests.Session()
    enriched: list[dict[str, str]] = []
    manual: list[dict[str, str]] = []

    for batch_start in range(0, len(selected), batch_size):
        for source in selected[batch_start : batch_start + batch_size]:
            row = {field: clean_text(source.get(field)) for field in REQUIRED_INPUT_COLUMNS}
            name_key = normalize_company_name(row["Company"])
            notes: list[str] = []
            source_url = ""
            careers_url = canonicalize_url(row["Careers URL"])

            if not name_key:
                status = "manual_review"
                detection = detect_ats("")
                notes.append("Company name is missing")
            elif counts[name_key] > 1:
                status = "duplicate_company"
                detection = detect_ats(careers_url)
                notes.append("Duplicate normalized company name in workbook")
            elif name_key in existing:
                entry = existing[name_key]
                detection = detection_from_registry(entry)
                source_url = registry_source_url(entry)
                if not careers_url:
                    careers_url = source_url
                status = "existing_config"
                notes.append("Preserved existing validated configuration")
                if row["Careers URL"] and canonicalize_url(row["Careers URL"]) != source_url:
                    notes.append("Workbook URL did not replace the existing configuration")
            elif name_key in reviewed_research:
                research = reviewed_research[name_key]
                research_valid, research_note = validate_research_row(research)
                researched_platform = clean_text(research.get("Detected Platform"))
                job_search_detection = detect_ats(research.get("Official Job Search URL"))
                detection = (
                    job_search_detection
                    if job_search_detection.platform == researched_platform
                    else ATSDetection(
                        platform=researched_platform,
                        slug=clean_text(research.get("ATS Slug")),
                        tenant=clean_text(research.get("Workday Tenant")),
                        wd_subdomain=clean_text(research.get("Workday Subdomain")),
                        site_id=clean_text(research.get("Workday Site ID")),
                    )
                )
                careers_url = clean_text(research.get("Official Careers URL"))
                source_url = clean_text(research.get("Source URL"))
                status = clean_text(research.get("Research Status")) if research_valid else "invalid"
                notes.extend(
                    [
                        research_note,
                        f"Official job search: {clean_text(research.get('Official Job Search URL'))}",
                        f"Required adapter: {clean_text(research.get('Required Adapter'))}",
                        (
                            "Existing adapter reusable: yes"
                            if parse_bool(research.get("Existing Adapter Reusable"))
                            else "Existing adapter reusable: no"
                        ),
                        clean_text(research.get("Research Notes")),
                    ]
                )
            elif not careers_url:
                detection = detect_ats("")
                status = "manual_review"
                notes.append("Missing careers URL; official-site research required")
            else:
                detection = detect_ats(careers_url)
                if validate_urls:
                    submitted_url = careers_url
                    validation = validate_public_https_url(careers_url, http)
                    source_url = submitted_url
                    careers_url = validation.final_url
                    notes.append(validation.note)
                    redirected_detection = detect_ats(validation.final_url)
                    redirected_valid, redirected_note = validate_detection_fields(redirected_detection)
                    notes.append(redirected_note)
                    if validation.valid and redirected_valid:
                        detection = redirected_detection
                        status = "validated"
                    else:
                        status = "manual_review"
                else:
                    source_url = careers_url
                    fields_valid, fields_note = validate_detection_fields(detection)
                    notes.append(fields_note)
                    if fields_valid:
                        status = "detected_unvalidated"
                        notes.append("Live URL validation was not requested")
                    else:
                        status = "manual_review"

            output = {
                **row,
                "Careers URL": careers_url,
                "ATS Platform": detection.platform,
                "ATS Slug": detection.slug,
                "Workday Tenant": detection.tenant,
                "Workday Subdomain": detection.wd_subdomain,
                "Workday Site ID": detection.site_id,
                "Import Status": status,
                "Validation Notes": "; ".join(dict.fromkeys(note for note in notes if note)),
                "Source URL": source_url,
            }
            enriched.append(output)
            if status not in {"existing_config", "validated", "validated_supported"}:
                manual.append(output)
    return enriched, manual


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--batch-size", type=int, default=25)
    parser.add_argument("--start-row", type=int, default=1, help="1-based data-row offset (header excluded)")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--research-file", type=Path, help="Reviewed official-URL research CSV")
    parser.add_argument("--dry-run", action="store_true", help="Report decisions without writing files or requesting URLs")
    return parser


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    rows = read_workbook(args.input)
    _, registry_entries = load_registry()
    research_rows = read_research_csv(args.research_file) if args.research_file else []
    effective_limit = args.limit if args.limit is not None else args.batch_size
    enriched, manual = enrich_rows(
        rows,
        registry_entries,
        start_row=args.start_row,
        limit=effective_limit,
        batch_size=args.batch_size,
        validate_urls=not args.dry_run,
        research_rows=research_rows,
    )
    summary = {
        "input": str(args.input),
        "workbook_rows": len(rows),
        "selected_rows": len(enriched),
        "start_row": args.start_row,
        "limit": effective_limit,
        "batch_size": args.batch_size,
        "dry_run": args.dry_run,
        "research_file": str(args.research_file) if args.research_file else None,
        "status_counts": dict(sorted(Counter(row["Import Status"] for row in enriched).items())),
        "companies": [
            {
                "rank": row["Rank"],
                "company": row["Company"],
                "status": row["Import Status"],
                "platform": row["ATS Platform"],
                "notes": row["Validation Notes"],
            }
            for row in enriched
        ],
    }
    if not args.dry_run:
        write_csv(args.output, enriched, ENRICHED_FIELDS)
        write_csv(DEFAULT_MANUAL_REVIEW, manual, ENRICHED_FIELDS)
        summary["output"] = str(args.output)
        summary["manual_review_output"] = str(DEFAULT_MANUAL_REVIEW)
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
