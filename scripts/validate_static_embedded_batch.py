"""Build and validate the bounded static-HTML/embedded-data rollout.

This module imports only registry and scanner code.  It never initializes
persistence, dashboard publishing, Supabase, Google Sheets, or SQLite.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Mapping, Optional

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import validate_company_registry as validator
from src.scraper.company_careers import scanner
from src.scraper.company_careers.adapters import embedded_data, static_html
from src.scraper.company_careers.registry import load_target_companies


PRODUCTION = ROOT / "config" / "target_companies.yaml"
PARENT_MAPPINGS = ROOT / "config" / "company_parent_source_mappings.yaml"
MASTER_WORKBOOK = ROOT / "config" / "company_watchlist.xlsx"
REMAINING = ROOT / "config" / "company_remaining_inventory_v2.csv"
V2_RESEARCH = ROOT / "config" / "company_platform_fingerprints_v2.csv"
V3_RESEARCH = ROOT / "config" / "company_unknown_investigation_v3.csv"
CANDIDATE_INVENTORY = ROOT / "config" / "company_static_embedded_candidates.csv"
CANDIDATE_REGISTRY = ROOT / "config" / "target_companies.static_embedded_candidates.yaml"
VALIDATED_REGISTRY = ROOT / "config" / "target_companies.static_embedded_validated.yaml"
REPORT = ROOT / "output" / "static_embedded_validation.json"

MAX_BATCH_SIZE = 5
MAX_PAGES = 2
MAX_JOBS = 20
MAX_SCRIPTS = 20
MAX_RESPONSE_BYTES = 3_000_000
MAX_SCRIPT_BYTES = 1_000_000
SAFE_STATUSES = {validator.VALIDATED_WITH_JOBS, validator.VALIDATED_ZERO_JOBS}
FAMILIES = {"static_html", "javascript_embedded_data"}


def _base(rank: int, name: str, platform: str, listing_url: str, endpoint_variant: str,
          allowed_hosts: list[str], **contract: Any) -> dict[str, Any]:
    entry: dict[str, Any] = {
        "name": name, "platform": platform, "careers_url": listing_url,
        "listing_url": listing_url, "endpoint_variant": endpoint_variant,
        "original_workbook_rank": rank, "allowed_apply_hosts": allowed_hosts,
        "max_pages": min(int(contract.pop("max_pages", 1)), MAX_PAGES),
        "max_jobs": MAX_JOBS, "max_scripts": MAX_SCRIPTS,
        "max_response_bytes": MAX_RESPONSE_BYTES, "max_script_bytes": MAX_SCRIPT_BYTES,
        "evidence_source": contract.pop("evidence_source", "config/company_remaining_inventory_v2.csv"),
        "evidence_confidence": "medium",
        "evidence_timestamp": contract.pop("evidence_timestamp", "2026-08-05T01:24:24+00:00"),
    }
    entry.update(contract)
    return entry


def _static(rank: int, name: str, url: str, hosts: list[str], *, card: str, title: str,
            link: str, id_regex: Optional[str] = None, id_attr: Optional[str] = None,
            listing: Optional[str] = None, location: Optional[str] = None,
            department: Optional[str] = None, pagination: Optional[str] = None,
            next_selector: Optional[str] = None, allow_fallback: bool = False,
            variant: str = "server_rendered_cards", **extra: Any) -> dict[str, Any]:
    return _base(
        rank, name, "static_html", url, variant, hosts,
        listing_selector=listing or card, job_card_selector=card, title_selector=title,
        detail_link_selector=link, external_id_regex=id_regex,
        external_id_attribute=id_attr, location_selector=location,
        department_selector=department, pagination_url_template=pagination,
        next_page_selector=next_selector, allow_url_id_fallback=allow_fallback, **extra,
    )


def _unsupported(rank: int, name: str, family: str, url: str, host: str, variant: str,
                 reason: str, *, stale: bool = False, source: str = "config/company_unknown_investigation_v3.csv",
                 timestamp: str = "2026-08-05T02:55:10+00:00") -> dict[str, Any]:
    common = dict(expected_validation_status="stale_source" if stale else "unsupported_variant",
                  expected_exclusion_reason=reason, evidence_source=source,
                  evidence_timestamp=timestamp)
    if family == "static_html":
        return _static(rank, name, url, [host], card="a[href]", title=":self", link=":self",
                       allow_fallback=True, variant=variant, **common)
    return _base(rank, name, family, url, variant, [host], embedded_script_selector="script#__NEXT_DATA__",
                 embedded_json_path="unsupported.items", id_field="id", title_field="title",
                 url_field="url", **common)


CANDIDATE_ENTRIES: list[dict[str, Any]] = [
    _static(12, "Waymo", "https://careers.withwaymo.com/jobs/search", ["careers.withwaymo.com"], card="article.job-search-results-card-col", title="h3.job-search-results-card-title a", link="h3.job-search-results-card-title a", id_regex=r"/jobs/([^/?#]+)", location=".job-search-results-card-location", pagination="https://careers.withwaymo.com/jobs/search?page={page}", allow_fallback=True),
    _unsupported(26, "ASML", "javascript_embedded_data", "https://www.asml.com/en/careers", "www.asml.com", "nextjs_landing_without_job_collection", "NEXT_DATA contains CMS and dictionary state but no reviewed jobs collection"),
    _static(27, "Applied Materials", "https://jobs.appliedmaterials.com/en/search-jobs", ["jobs.appliedmaterials.com"], card="li.search-results-list__item", title="h2.search-results-list__job-title a", link="h2.search-results-list__job-title a", id_regex=r"/(\d+)(?:[?#].*)?$", location=".search-results-list__location", pagination="https://jobs.appliedmaterials.com/en/search-jobs?p={page}"),
    _unsupported(62, "Adyen", "javascript_embedded_data", "https://careers.adyen.com/", "careers.adyen.com", "nuxt_reference_state", "Nuxt reference-encoded state has no reviewed deterministic job path"),
    _static(70, "Ripple", "https://ripple.com/careers/", ["ripple.com"], card="a[href*='/careers/all-jobs/job/']", title=":self", link=":self", id_regex=r"/job/(\d+)", variant="server_rendered_links", evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:53:31+00:00"),
    _base(79, "DRW", "javascript_embedded_data", "https://www.drw.com/work-at-drw/listings", "next_data", ["www.drw.com"], embedded_script_selector="script#__NEXT_DATA__", embedded_json_path="props.pageProps.jobData.en", id_field="id", title_field="title", location_field="locations", url_field="slug", department_field="career_categories", apply_url_template="https://www.drw.com/work-at-drw/listings/{slug}", evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:53:39+00:00"),
    _unsupported(80, "IMC Trading", "static_html", "https://www.imc.com/us/search-careers", "www.imc.com", "ambiguous_link_cards", "Job and apply anchors share unstable generated containers; title/location contract is ambiguous", stale=True, timestamp="2026-08-05T02:53:42+00:00"),
    _unsupported(100, "Yaskawa Electric", "static_html", "https://www.yaskawa-global.com/", "www.yaskawa-global.com", "regional_landing_links", "Page links to regional careers sites rather than deterministic job records", stale=True, timestamp="2026-08-05T02:53:53+00:00"),
    _static(115, "General Motors", "https://search-careers.gm.com/en/jobs/", ["search-careers.gm.com"], card="div.card-job", title="h2.card-title a", link="h2.card-title a", id_regex=r"/jobs/(jr-[0-9]+)/", location="ul.job-meta li:nth-child(2)", department="ul.job-meta li:first-child", pagination="https://search-careers.gm.com/en/jobs/?page={page}"),
    _unsupported(126, "Aptiv", "static_html", "https://www.aptiv.com/en/jobs", "www.aptiv.com", "careers_landing_without_jobs", "Current page is a careers landing page and does not expose job records", timestamp="2026-08-05T02:54:54+00:00"),
    _static(130, "Mobileye", "https://careers.mobileye.com/jobs", ["careers.mobileye.com"], card="div.jobItem", title="h3", link="a[href^='/jobs/']", id_regex=r"/([0-9a-f]{8}-[0-9a-f-]{27,})$", location=".location", department=".department", max_response_bytes=1_500_000),
    _unsupported(131, "Aurora Innovation", "javascript_embedded_data", "https://aurora.tech/careers", "aurora.tech", "nuxt_reference_state", "Nuxt reference-encoded state has no reviewed deterministic job path"),
    _unsupported(145, "Dassault Systèmes", "javascript_embedded_data", "https://www.3ds.com/careers", "www.3ds.com", "nuxt_reference_state", "Nuxt state contains a careers landing page, not a reviewed job collection", timestamp="2026-08-05T02:55:47+00:00"),
    _static(151, "Skydio", "https://www.skydio.com/ja-jp/careers", ["www.skydio.com"], card="div.job-listings__job", title="h5.job-listings__title", link="a[href*='/jobs/']", id_regex=r"/jobs/([0-9a-f-]{36})", location="p.job-listings__location", evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:55:59+00:00"),
    _unsupported(154, "Figure AI", "javascript_embedded_data", "https://www.figure.ai/careers", "www.figure.ai", "nextjs_thin_shell", "NEXT_DATA has navigation content but no reviewed job collection", timestamp="2026-08-05T02:56:00+00:00"),
    _static(156, "Unitree Robotics", "https://www.unitree.com/position/", ["www.unitree.com"], card="ul.posts li", title="h3", link="a[href^='/position/']", id_regex=r"/position/(\d+)", location=".location", department=".category", evidence_timestamp="2026-08-05T01:28:26+00:00"),
    _static(161, "Caterpillar", "https://careers.caterpillar.com/en/jobs/", ["careers.caterpillar.com"], card="div.card-job", title="h2.card-title a", link="h2.card-title a", id_attr="data-id", location="ul.job-meta li:first-child", pagination="https://careers.caterpillar.com/en/jobs/?page={page}"),
    _static(162, "Intuitive Surgical", "https://careers.intuitive.com/en/jobs/", ["careers.intuitive.com"], card="div.card-job", title="h2.card-title a", link="h2.card-title a", id_attr="data-id", location="ul.job-meta li:first-child", department="ul.job-meta li:nth-child(2)", pagination="https://careers.intuitive.com/en/jobs/?page={page}"),
    _unsupported(181, "Tourmaline Oil", "javascript_embedded_data", "https://www.tourmalineoil.com/careers", "www.tourmalineoil.com", "nextjs_document_link_only", "Page exposes a PDF posting but no stable structured job collection", timestamp="2026-08-05T02:56:22+00:00"),
    _static(202, "Coupang", "https://www.coupang.jobs/en/jobs/", ["www.coupang.jobs"], card="div.card-job", title="h2.card-title a", link="h2.card-title a", id_regex=r"(?:gh_jid=|/jobs/)(\d+)", location="ul.job-meta li:first-child", pagination="https://www.coupang.jobs/en/jobs/?page={page}"),
    _static(205, "Etsy", "https://careers.etsy.com/jobs/search", ["careers.etsy.com"], card="article.job-search-results-card-col", title="h3.job-search-results-card-title a", link="h3.job-search-results-card-title a", id_regex=r"/jobs/([^/?#]+)", location=".job-search-results-card-location", pagination="https://careers.etsy.com/jobs/search?page={page}", allow_fallback=True, evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:56:49+00:00"),
    _static(210, "Zalando", "https://jobs.zalando.com/en/jobs", ["jobs.zalando.com"], card="main a[href^='/en/jobs/']", title=":self", link=":self", id_regex=r"/jobs/(\d+)-", variant="server_rendered_links", evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:56:57+00:00"),
    _static(219, "Flexport", "https://www.flexport.com/careers/jobs/", ["job-boards.greenhouse.io"], card="a[href*='job-boards.greenhouse.io/flexport/jobs/']", title=":self", link=":self", id_regex=r"/jobs/(\d+)", variant="server_rendered_links"),
    _unsupported(226, "LinkedIn", "static_html", "https://careers.linkedin.com/", "careers.linkedin.com", "external_session_job_search", "Search Jobs leaves the official careers host for a session-oriented LinkedIn company page", stale=True, source="config/company_platform_fingerprints_v2.csv", timestamp="2026-08-05T01:29:19+00:00"),
    _unsupported(237, "LINE", "static_html", "https://careers.linecorp.com/jobs", "careers.linecorp.com", "empty_shell_without_contract", "HTTP 200 page has no server-rendered job cards or conclusive empty-list marker", source="config/company_platform_fingerprints_v2.csv", timestamp="2026-08-05T01:29:19+00:00"),
    _unsupported(243, "Patreon", "javascript_embedded_data", "https://www.patreon.com/careers", "www.patreon.com", "nextjs_cms_without_job_collection", "NEXT_DATA contains CMS page content but no reviewed deterministic job path", timestamp="2026-08-05T02:57:10+00:00"),
    _base(246, "Bluesky", "javascript_embedded_data", "https://bsky.social/about/join", "next_data", ["jobs.gem.com"], embedded_script_selector="script#__NEXT_DATA__", embedded_json_path="props.pageProps.jobs", id_field="id", title_field="title", location_field="location", url_field="applyUrl", date_field="updatedAt", department_field="department", remote_field="locationType", evidence_source="config/company_platform_fingerprints_v2.csv", evidence_timestamp="2026-08-05T01:29:46+00:00"),
]

EXPECTED_COMPANIES = [entry["name"] for entry in CANDIDATE_ENTRIES]
STATIC_COMPANIES = [entry["name"] for entry in CANDIDATE_ENTRIES if entry["platform"] == "static_html"]
EMBEDDED_COMPANIES = [entry["name"] for entry in CANDIDATE_ENTRIES if entry["platform"] == "javascript_embedded_data"]


def _rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _master_ranks() -> dict[str, int]:
    workbook = openpyxl.load_workbook(MASTER_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    result = {str(row[headers.index("Company")]).strip(): int(row[headers.index("Rank")])
              for row in rows if row[headers.index("Company")] not in (None, "")}
    workbook.close()
    return result


def _derived_research_candidates() -> list[str]:
    remaining = {row["company"]: row for row in _rows(REMAINING)}
    names = [row["company"] for row in _rows(V2_RESEARCH)
             if row.get("platform_family", "").casefold() in FAMILIES
             and remaining.get(row["company"], {}).get("currently_uncovered", "").casefold() == "true"]
    names.extend(row["company"] for row in _rows(V3_RESEARCH)
                 if row.get("platform_family", "").casefold() in FAMILIES)
    ranks = _master_ranks()
    return sorted(dict.fromkeys(names), key=ranks.__getitem__)


def reconcile_candidates() -> None:
    derived = _derived_research_candidates()
    if derived != EXPECTED_COMPANIES:
        raise ValueError(f"Research artifacts do not derive the exact candidate set: expected={EXPECTED_COMPANIES!r}, found={derived!r}")
    ranks = _master_ranks()
    configured = [int(entry["original_workbook_rank"]) for entry in CANDIDATE_ENTRIES]
    if configured != [ranks[name] for name in EXPECTED_COMPANIES] or configured != sorted(configured):
        raise ValueError("Candidate ordering does not match the original master workbook")
    _, production = validator.load_registry(PRODUCTION)
    production_names = {validator._name_key(entry) for entry in production}
    mappings = yaml.safe_load(PARENT_MAPPINGS.read_text(encoding="utf-8")) or {}
    mapped = {" ".join(str(item.get("child_company") or "").casefold().split())
              for item in mappings.get("company_parent_source_mappings", []) if item.get("enabled") is True}
    conflicts = [entry["name"] for entry in CANDIDATE_ENTRIES if validator._name_key(entry) in production_names | mapped]
    if conflicts:
        raise ValueError(f"Candidates already covered: {conflicts}")
    if validator._duplicate_company_owners(production, CANDIDATE_ENTRIES) or validator._duplicate_owners(production, CANDIDATE_ENTRIES):
        raise ValueError("Candidate companies or source identities are duplicated")


def build_candidate_registry() -> None:
    reconcile_candidates()
    content = validator.build_candidate_registry_bytes(PRODUCTION, CANDIDATE_ENTRIES)
    if not content.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Candidate registry does not preserve the production byte prefix")
    validator._write_bytes_if_changed(CANDIDATE_REGISTRY, content)
    if len(load_target_companies(str(CANDIDATE_REGISTRY))) != len(load_target_companies(str(PRODUCTION))) + 27:
        raise ValueError("Candidate registry is not loader-compatible")


def _status(entry: Mapping[str, Any], raw: Mapping[str, Any]) -> str:
    expected = entry.get("expected_validation_status")
    if expected:
        return str(expected)
    status = str(raw.get("validation_status") or "failed_technical")
    error = str(raw.get("error") or "").casefold()
    if raw.get("http_status") == 202 or "http 202" in error:
        return "temporary_failure"
    if status == validator.MALFORMED_RESPONSE:
        if "stable identifier" in error:
            return "missing_stable_id"
        if "allowed-host" in error or "apply url" in error:
            return "invalid_apply_url"
        return "malformed_response"
    if status == validator.INVALID_CONFIGURATION:
        return "invalid_configuration"
    if status == validator.UNSUPPORTED_RESPONSE:
        return "failed_technical"
    if status == validator.TEMPORARY_FAILURE:
        return "temporary_failure"
    return status


def _report_row(entry: Mapping[str, Any], raw: Mapping[str, Any]) -> dict[str, Any]:
    status = _status(entry, raw)
    safe = status in SAFE_STATUSES
    pagination = raw.get("pagination_result") if isinstance(raw.get("pagination_result"), Mapping) else {}
    diagnostics = list(pagination.get("diagnostics") or [])
    reason = entry.get("expected_exclusion_reason") or raw.get("error")
    return {
        "company": entry["name"], "master_rank": entry["original_workbook_rank"],
        "expected_collector_family": entry["platform"], "endpoint_variant": entry["endpoint_variant"],
        "pages_requested": int(raw.get("pages_requested") or 0),
        "scripts_inspected": int(pagination.get("scripts_inspected") or 0),
        "bytes_inspected": int(pagination.get("bytes_inspected") or 0),
        "jobs_observed": int(raw.get("jobs_returned") or 0),
        "normalized_jobs_inspected": int(raw.get("jobs_normalized") or 0),
        "stable_id_result": raw.get("stable_job_id_result"),
        "apply_host_result": raw.get("apply_url_result"),
        "pagination_result": dict(pagination),
        "repeated_page_result": bool(pagination.get("repeated_page")),
        "valid_empty_result": bool(pagination.get("zero_job_contract")),
        "browser_check": "not_required" if safe else ("failed" if status == "browser_required" else "not_observed"),
        "session_check": "not_required" if safe else ("failed" if status == "session_required" else "not_observed"),
        "authentication_check": "not_required" if safe else ("failed" if status == "authentication_required" else "not_observed"),
        "blocked_response_check": "not_blocked" if safe else ("failed" if status == "blocked" else "not_observed"),
        "duplicate_company_result": raw.get("duplicate_company_status") or "unique",
        "duplicate_source_result": raw.get("duplicate_source_status") or "unique",
        "status": status, "safe_to_activate": safe,
        "warnings": [str(raw["warning"])] if raw.get("warning") else [],
        "diagnostics": diagnostics + ([str(reason)] if reason else []),
    }


def _write_inventory(rows: list[Mapping[str, Any]]) -> None:
    by_name = {row["company"]: row for row in rows}
    fields = ["master_rank", "company", "expected_platform", "endpoint_variant", "careers_url", "listing_url",
              "pagination_contract", "identifier_contract", "apply_url_contract", "allowed_apply_hosts",
              "listing_selectors_or_json_path", "evidence_source", "evidence_confidence", "evidence_timestamp",
              "prior_classification", "validation_status", "safe_to_activate", "warnings"]
    CANDIDATE_INVENTORY.parent.mkdir(parents=True, exist_ok=True)
    with CANDIDATE_INVENTORY.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for entry in CANDIDATE_ENTRIES:
            report = by_name.get(entry["name"], {})
            writer.writerow({
                "master_rank": entry["original_workbook_rank"], "company": entry["name"],
                "expected_platform": entry["platform"], "endpoint_variant": entry["endpoint_variant"],
                "careers_url": entry["careers_url"], "listing_url": entry["listing_url"],
                "pagination_contract": entry.get("pagination_url_template") or entry.get("next_page_selector") or "single_public_page",
                "identifier_contract": entry.get("external_id_attribute") or entry.get("external_id_regex") or entry.get("id_field") or "unsupported",
                "apply_url_contract": entry.get("url_field") or entry.get("detail_link_selector") or "unsupported",
                "allowed_apply_hosts": ";".join(entry["allowed_apply_hosts"]),
                "listing_selectors_or_json_path": entry.get("job_card_selector") or entry.get("embedded_json_path") or "",
                "evidence_source": entry["evidence_source"], "evidence_confidence": entry["evidence_confidence"],
                "evidence_timestamp": entry["evidence_timestamp"], "prior_classification": entry["platform"],
                "validation_status": report.get("status", "not_run"),
                "safe_to_activate": str(bool(report.get("safe_to_activate"))).lower(),
                "warnings": "; ".join(
                    item if isinstance(item, str) else json.dumps(item, ensure_ascii=False, sort_keys=True)
                    for item in (report.get("diagnostics") or [])
                ),
            })


def run(*, live: bool = True, timeout: int = 25) -> dict[str, Any]:
    production_before = PRODUCTION.read_bytes()
    build_candidate_registry()
    raw_results: list[dict[str, Any]] = []
    if live:
        _, production_entries = validator.load_registry(PRODUCTION)
        duplicate_sources = validator._duplicate_owners(production_entries, CANDIDATE_ENTRIES)
        duplicate_names = validator._duplicate_company_owners(production_entries, CANDIDATE_ENTRIES)
        for start in range(0, len(CANDIDATE_ENTRIES), MAX_BATCH_SIZE):
            for offset, entry in enumerate(CANDIDATE_ENTRIES[start:start + MAX_BATCH_SIZE]):
                index = start + offset
                raw_results.append(validator.validate_company(
                    entry, duplicate_of=duplicate_sources.get(index), duplicate_company_of=duplicate_names.get(index),
                    timeout=timeout, attempts=1, retry_delay=0, sleep=lambda _: None,
                    adapters=scanner.ADAPTERS, max_normalized_jobs=MAX_JOBS,
                ))
    else:
        raw_results = [{"company": entry["name"], "validation_status": "not_run"} for entry in CANDIDATE_ENTRIES]
    rows = [_report_row(entry, raw) for entry, raw in zip(CANDIDATE_ENTRIES, raw_results)]
    safe_names = {row["company"] for row in rows if row["safe_to_activate"]}
    safe_entries = [entry for entry in CANDIDATE_ENTRIES if entry["name"] in safe_names]
    validated = validator.build_candidate_registry_bytes(PRODUCTION, safe_entries)
    if not validated.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Validated registry does not preserve production prefix")
    validator._write_bytes_if_changed(VALIDATED_REGISTRY, validated)
    _write_inventory(rows)
    counts = Counter(row["status"] for row in rows)
    production_count = len(load_target_companies(str(PRODUCTION)))
    summary = {
        "reviewed_candidates": len(rows), "static_html_candidates": len(STATIC_COMPANIES),
        "embedded_data_candidates": len(EMBEDDED_COMPANIES),
        "validated_with_jobs": counts[validator.VALIDATED_WITH_JOBS],
        "validated_with_zero_jobs": counts[validator.VALIDATED_ZERO_JOBS],
        "excluded": sum(not row["safe_to_activate"] and row["status"] != "temporary_failure" for row in rows),
        "temporary_failures": counts["temporary_failure"],
        "candidate_registry_count": production_count + len(CANDIDATE_ENTRIES),
        "validated_registry_count": production_count + len(safe_entries),
        "projected_original_company_coverage": 99 + len(safe_entries),
    }
    report = {"summary": summary, "results": rows}
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if PRODUCTION.read_bytes() != production_before:
        raise ValueError("Production registry changed during validation")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-live", action="store_true", help="Build artifacts without HTTP validation")
    parser.add_argument("--timeout", type=int, default=25)
    args = parser.parse_args()
    report = run(live=not args.no_live, timeout=max(1, args.timeout))
    print(json.dumps(report["summary"], indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
