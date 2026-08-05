"""Build and validate the deterministic Avature/iCIMS rollout candidates.

The command imports only registry/scanner code.  It does not initialize a
database, Supabase, Google Sheets, or the dashboard publisher.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional
from urllib.parse import urlparse

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import validate_company_registry as validator
from src.scraper.company_careers import scanner
from src.scraper.company_careers.registry import load_target_companies


PRODUCTION = ROOT / "config" / "target_companies.yaml"
PARENT_MAPPINGS = ROOT / "config" / "company_parent_source_mappings.yaml"
MASTER_WORKBOOK = ROOT / "config" / "company_watchlist.xlsx"
V2_RESEARCH = ROOT / "config" / "company_remaining_inventory_v2.csv"
V3_RESEARCH = ROOT / "config" / "company_unknown_investigation_v3.csv"
CANDIDATE_INVENTORY = ROOT / "config" / "company_avature_icims_candidates.csv"
CANDIDATE_REGISTRY = ROOT / "config" / "target_companies.avature_icims_candidates.yaml"
VALIDATED_REGISTRY = ROOT / "config" / "target_companies.avature_icims_validated.yaml"
REPORT = ROOT / "output" / "avature_icims_validation.json"

MAX_BATCH_SIZE = 5
MAX_PAGES = 2
MAX_JOBS = 20
SAFE_STATUSES = {validator.VALIDATED_WITH_JOBS, validator.VALIDATED_ZERO_JOBS}
INDUSTRY_NOT_RECORDED = "Not recorded in repository master workbook"


def _entry(
    rank: int,
    name: str,
    platform: str,
    careers_url: str,
    search_url: str,
    endpoint_variant: str,
    *,
    tenant: str,
    board_id: str,
    allowed_apply_hosts: list[str],
    endpoint_url: Optional[str] = None,
    site_id: Optional[str] = None,
    page_size: int = 10,
    evidence_source: str = "config/company_remaining_inventory_v2.csv",
    evidence_timestamp: str = "2026-08-05T01:24:24+00:00",
    expected_validation_status: Optional[str] = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "name": name,
        "platform": platform,
        "tenant": tenant,
        "board_id": board_id,
        "careers_url": careers_url,
        "search_url": search_url,
        "endpoint_variant": endpoint_variant,
        "original_workbook_rank": rank,
        "industry": INDUSTRY_NOT_RECORDED,
        "allowed_apply_hosts": allowed_apply_hosts,
        "page_size": page_size,
        "max_pages": MAX_PAGES,
        "max_jobs": MAX_JOBS,
        "evidence_source": evidence_source,
        "evidence_confidence": "high",
        "evidence_timestamp": evidence_timestamp,
    }
    if endpoint_url:
        result["endpoint_url"] = endpoint_url
    if site_id:
        result["site_id"] = site_id
    if expected_validation_status:
        result["expected_validation_status"] = expected_validation_status
    return result


CANDIDATE_ENTRIES: list[dict[str, Any]] = [
    _entry(17, "AMD", "icims", "https://www.amd.com/en/corporate/careers.html", "https://careers.amd.com/careers-home/jobs", "icims_jibe_api", tenant="amd", board_id="amd-public", endpoint_url="https://careers.amd.com/api/jobs", allowed_apply_hosts=["careers-amd.icims.com", "canadacareers-amd.icims.com", "global-external-amd.icims.com"]),
    _entry(20, "TSMC", "avature", "https://www.tsmc.com/static/english/careers/index.htm", "https://careers.tsmc.com/en_US/careers/SearchJobs", "avature_portal_html", tenant="careers.tsmc.com", board_id="6", site_id="en_US/careers", allowed_apply_hosts=["careers.tsmc.com"]),
    _entry(23, "ARM", "icims", "https://careers.arm.com/", "https://careers.arm.com/en/search-jobs", "icims_radancy_html", tenant="arm", board_id="33099", allowed_apply_hosts=["careers.arm.com"], page_size=15),
    _entry(24, "Synopsys", "avature", "https://www.synopsys.com/careers/global.html", "https://careers.synopsys.com/search-jobs", "avature_radancy_html", tenant="synopsys.avature.net", board_id="44408", allowed_apply_hosts=["careers.synopsys.com"], page_size=15),
    _entry(40, "Lattice Semiconductor", "icims", "https://careers-latticesemi.icims.com/jobs/intro?hashed=-625919477", "https://careers-latticesemi.icims.com/jobs/search?hashed=-625919477", "icims_legacy_portal_html", tenant="latticesemi", board_id="latticesemi-public", site_id="-625919477", allowed_apply_hosts=["careers-latticesemi.icims.com"], page_size=50, evidence_source="config/company_unknown_investigation_v3.csv", evidence_timestamp="2026-08-05T02:53:16+00:00"),
    _entry(44, "GitHub", "icims", "https://www.github.careers/careers-home", "https://www.github.careers/careers-home/jobs", "icims_jibe_api", tenant="githubinc", board_id="githubinc-public", endpoint_url="https://www.github.careers/api/jobs", allowed_apply_hosts=["careers-githubinc.icims.com", "globalcareers-githubinc.icims.com"]),
    _entry(64, "Intuit", "avature", "https://www.intuit.com/careers/", "https://jobs.intuit.com/search-jobs", "avature_radancy_html", tenant="intuit.avature.net", board_id="27595", allowed_apply_hosts=["jobs.intuit.com"], page_size=15),
    _entry(77, "Two Sigma", "avature", "https://www.twosigma.com/careers/", "https://careers.twosigma.com/careers", "avature_portal_html", tenant="careers.twosigma.com", board_id="106", site_id="careers", allowed_apply_hosts=["careers.twosigma.com"], page_size=4),
    _entry(82, "SIG", "icims", "https://sig.com/careers/", "https://careers.sig.com/jobs", "icims_jibe_api", tenant="sig", board_id="sig-public", endpoint_url="https://careers.sig.com/api/jobs", allowed_apply_hosts=["careers-sig.icims.com"]),
    _entry(87, "Bank of America", "avature", "https://careers.bankofamerica.com/", "https://careers.bankofamerica.com/en-us/job-search.html", "avature_aem_html", tenant="bankofamerica.avature.net", board_id="professional-us", allowed_apply_hosts=["careers.bankofamerica.com"]),
    _entry(91, "Siemens", "avature", "https://www.siemens.com/global/en/company/jobs.html", "https://jobs.siemens.com/en_US/externaljobs/SearchJobs", "avature_folder_portal", tenant="jobs.siemens.com", board_id="144", site_id="en_US/externaljobs", allowed_apply_hosts=["jobs.siemens.com"], page_size=6, expected_validation_status="unsupported_variant"),
    _entry(116, "Rivian", "icims", "https://careers.rivian.com/", "https://careers.rivian.com/careers-home/jobs", "icims_jibe_api", tenant="rivian", board_id="rivian-public", endpoint_url="https://careers.rivian.com/api/jobs", allowed_apply_hosts=["us-careers-rivian.icims.com", "sb-careers-rivian.icims.com"]),
    _entry(148, "Lenovo", "avature", "https://jobs.lenovo.com/", "https://jobs.lenovo.com/en_US/careers/SearchJobs", "avature_portal_html", tenant="jobs.lenovo.com", board_id="4", site_id="en_US/careers", allowed_apply_hosts=["jobs.lenovo.com"]),
    _entry(172, "bp", "avature", "https://www.bp.com/en/global/corporate/careers.html", "https://careers.bp.com/", "stale_workday_migration", tenant="bp.avature.net", board_id="legacy-bp", allowed_apply_hosts=["careers.bp.com"], expected_validation_status="stale_source"),
    _entry(211, "Wayfair", "avature", "https://www.wayfair.com/careers", "https://www.wayfair.com/careers/jobs", "session_react_portal", tenant="wayfair.avature.net", board_id="legacy-wayfair", allowed_apply_hosts=["www.wayfair.com"], expected_validation_status="session_required"),
]

EXPECTED_COMPANIES = [entry["name"] for entry in CANDIDATE_ENTRIES]


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _derived_research_candidates() -> list[str]:
    rows = _csv_rows(V2_RESEARCH)
    names = [
        row["company"]
        for row in rows
        if row.get("currently_uncovered", "").casefold() == "true"
        and row.get("current_platform_family", "").casefold() in {"avature", "icims"}
        and row.get("stale_fingerprint", "").casefold() != "true"
    ]
    for row in _csv_rows(V3_RESEARCH):
        if (
            row.get("platform_family", "").casefold() in {"avature", "icims"}
            and row.get("primary_classification", "").casefold() in {"avature", "icims"}
            and row.get("anonymous_access", "").casefold() == "true"
            and row.get("blocked", "").casefold() != "true"
            and row.get("browser_required", "").casefold() != "true"
        ):
            names.append(row["company"])
    rank_by_name = _master_ranks()
    return sorted(dict.fromkeys(names), key=rank_by_name.__getitem__)


def _master_ranks() -> dict[str, int]:
    workbook = openpyxl.load_workbook(MASTER_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    rank_index = headers.index("Rank")
    company_index = headers.index("Company")
    result = {
        str(row[company_index]).strip(): int(row[rank_index])
        for row in rows
        if row[company_index] not in (None, "")
    }
    workbook.close()
    return result


def reconcile_candidates() -> None:
    derived = _derived_research_candidates()
    if derived != EXPECTED_COMPANIES:
        raise ValueError(
            "Research artifacts do not derive the exact Avature/iCIMS candidate set: "
            f"expected={EXPECTED_COMPANIES!r}, found={derived!r}"
        )
    ranks = _master_ranks()
    actual_ranks = [ranks[entry["name"]] for entry in CANDIDATE_ENTRIES]
    configured_ranks = [entry["original_workbook_rank"] for entry in CANDIDATE_ENTRIES]
    if actual_ranks != configured_ranks or configured_ranks != sorted(configured_ranks):
        raise ValueError("Candidate ordering does not match the original 246-company workbook")

    _, production_entries = validator.load_registry(PRODUCTION)
    production_names = {validator._name_key(entry) for entry in production_entries}
    mappings = yaml.safe_load(PARENT_MAPPINGS.read_text(encoding="utf-8")) or {}
    mapped_names = {
        " ".join(str(item.get("child_company") or "").casefold().split())
        for item in mappings.get("company_parent_source_mappings", [])
        if item.get("enabled") is True
    }
    conflicts = [
        entry["name"]
        for entry in CANDIDATE_ENTRIES
        if validator._name_key(entry) in production_names | mapped_names
    ]
    if conflicts:
        raise ValueError(f"Candidates already covered in production: {conflicts}")
    duplicate_names = validator._duplicate_company_owners(production_entries, CANDIDATE_ENTRIES)
    duplicate_sources = validator._duplicate_owners(production_entries, CANDIDATE_ENTRIES)
    if duplicate_names or duplicate_sources:
        raise ValueError(
            f"Candidate duplicates: companies={duplicate_names}, sources={duplicate_sources}"
        )
    for entry in CANDIDATE_ENTRIES:
        if int(entry["max_pages"]) > MAX_PAGES or int(entry["max_jobs"]) > MAX_JOBS:
            raise ValueError(f"{entry['name']}: validation limits exceeded")


def build_candidate_registry() -> bytes:
    reconcile_candidates()
    content = validator.build_candidate_registry_bytes(PRODUCTION, CANDIDATE_ENTRIES)
    validator._write_bytes_if_changed(CANDIDATE_REGISTRY, content)
    _, production_entries = validator.load_registry(PRODUCTION)
    if not content.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Candidate registry does not preserve the production byte prefix")
    loaded = load_target_companies(str(CANDIDATE_REGISTRY))
    if len(loaded) != len(production_entries) + len(CANDIDATE_ENTRIES):
        raise ValueError("Production registry loader rejected candidate entries")
    return content


def _status(entry: Mapping[str, Any], raw: Mapping[str, Any]) -> str:
    expected = entry.get("expected_validation_status")
    if expected and raw.get("validation_status") not in SAFE_STATUSES:
        return str(expected)
    status = str(raw.get("validation_status") or "failed_technical")
    if status == validator.UNSUPPORTED_RESPONSE:
        return "malformed_response" if raw.get("http_status") == 200 else "failed_technical"
    if status == validator.TEMPORARY_FAILURE:
        return "temporary_failure"
    if status == validator.STALE_FINGERPRINT:
        return "stale_source"
    return status


def _report_result(entry: Mapping[str, Any], raw: Mapping[str, Any]) -> dict[str, Any]:
    status = _status(entry, raw)
    safe = status in SAFE_STATUSES
    pagination = raw.get("pagination_result")
    if not isinstance(pagination, Mapping):
        pagination = {
            "pages_fetched": int(raw.get("pages_requested") or 0),
            "maximum_pages": MAX_PAGES,
            "maximum_results": MAX_JOBS,
            "stop_reason": "configuration_rejected" if not safe else "unknown",
            "repeated_page": False,
            "zero_job_contract": False,
        }
    error = raw.get("error")
    warnings = [str(raw["warning"])] if raw.get("warning") else []
    return {
        "company": entry["name"],
        "master_rank": entry["original_workbook_rank"],
        "industry": entry["industry"],
        "platform": entry["platform"],
        "endpoint_identifiers": validator.configured_identifiers(entry),
        "endpoint_url": validator.endpoint_for(entry),
        "source_identifier": validator.source_key(entry),
        "pages_requested": int(raw.get("pages_requested") or 0),
        "jobs_observed": int(raw.get("jobs_returned") or 0),
        "normalized_jobs_inspected": int(raw.get("jobs_normalized") or 0),
        "stable_id_result": raw.get("stable_job_id_result"),
        "apply_url_result": raw.get("apply_url_result"),
        "allowed_host_result": raw.get("apply_url_result"),
        "pagination_result": dict(pagination),
        "repeated_page_result": bool(pagination.get("repeated_page")),
        "zero_job_contract_result": bool(pagination.get("zero_job_contract")),
        "duplicate_company_result": raw.get("duplicate_company_status"),
        "duplicate_source_result": raw.get("duplicate_source_status"),
        "authentication_check": "not_required" if safe else _failure_check(status, "authentication_required"),
        "browser_check": "not_required" if safe else _failure_check(status, "browser_required"),
        "session_check": "not_required" if safe else _failure_check(status, "session_required"),
        "blocked_response_check": "not_blocked" if safe else _failure_check(status, "blocked"),
        "warnings": warnings,
        "error": error,
        "status": status,
        "safe_to_activate": safe,
    }


def _failure_check(status: str, matching: str) -> str:
    return "detected" if status == matching else "not_observed"


def _inventory_rows(results: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    by_name = {str(result["company"]): result for result in results}
    rows: list[dict[str, Any]] = []
    for entry in CANDIDATE_ENTRIES:
        result = by_name.get(entry["name"], {})
        rows.append(
            {
                "master_rank": entry["original_workbook_rank"],
                "company": entry["name"],
                "industry": entry["industry"],
                "expected_platform": entry["platform"],
                "careers_url": entry["careers_url"],
                "search_url": entry["search_url"],
                "endpoint_url": entry.get("endpoint_url", ""),
                "tenant": entry.get("tenant", ""),
                "site_id": entry.get("site_id", ""),
                "board_id": entry.get("board_id", ""),
                "endpoint_variant": entry["endpoint_variant"],
                "allowed_apply_hosts": ";".join(entry["allowed_apply_hosts"]),
                "evidence_source": entry["evidence_source"],
                "evidence_confidence": entry["evidence_confidence"],
                "evidence_timestamp": entry["evidence_timestamp"],
                "validation_status": result.get("status", "pending"),
                "safe_to_activate": str(bool(result.get("safe_to_activate"))).lower(),
                "warnings": "; ".join(result.get("warnings") or []),
            }
        )
    return rows


def _write_inventory(results: Iterable[Mapping[str, Any]]) -> None:
    rows = _inventory_rows(results)
    fieldnames = list(rows[0])
    CANDIDATE_INVENTORY.parent.mkdir(parents=True, exist_ok=True)
    with CANDIDATE_INVENTORY.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def validate(
    *,
    timeout: int = validator.DEFAULT_TIMEOUT,
    batch_delay: float = validator.DEFAULT_BATCH_DELAY,
    sleep: Any = time.sleep,
    session_factory: Any = None,
) -> dict[str, Any]:
    build_candidate_registry()
    _, production_entries = validator.load_registry(PRODUCTION)
    duplicate_sources = validator._duplicate_owners(production_entries, CANDIDATE_ENTRIES)
    duplicate_companies = validator._duplicate_company_owners(production_entries, CANDIDATE_ENTRIES)
    raw_results: list[dict[str, Any]] = []
    factory = session_factory or __import__("requests").Session
    for batch_start in range(0, len(CANDIDATE_ENTRIES), MAX_BATCH_SIZE):
        batch = CANDIDATE_ENTRIES[batch_start : batch_start + MAX_BATCH_SIZE]
        for offset, entry in enumerate(batch):
            index = batch_start + offset
            raw_results.append(
                validator.validate_company(
                    entry,
                    duplicate_of=duplicate_sources.get(index),
                    duplicate_company_of=duplicate_companies.get(index),
                    timeout=timeout,
                    attempts=validator.DEFAULT_ATTEMPTS,
                    retry_delay=validator.DEFAULT_RETRY_DELAY,
                    session_factory=factory,
                    sleep=sleep,
                    adapters=scanner.ADAPTERS,
                    max_normalized_jobs=MAX_JOBS,
                )
            )
        if batch_start + MAX_BATCH_SIZE < len(CANDIDATE_ENTRIES) and batch_delay > 0:
            sleep(batch_delay)

    results = [
        _report_result(entry, raw)
        for entry, raw in zip(CANDIDATE_ENTRIES, raw_results)
    ]
    safe_names = {result["company"] for result in results if result["safe_to_activate"]}
    safe_entries = [entry for entry in CANDIDATE_ENTRIES if entry["name"] in safe_names]
    validated_bytes = validator.build_validated_registry_bytes(PRODUCTION, safe_entries)
    validator._write_bytes_if_changed(VALIDATED_REGISTRY, validated_bytes)
    if not validated_bytes.startswith(PRODUCTION.read_bytes()):
        raise ValueError("Validated registry does not preserve the production byte prefix")
    status_counts = Counter(str(result["status"]) for result in results)
    candidate_count = len(production_entries) + len(CANDIDATE_ENTRIES)
    validated_count = len(production_entries) + len(safe_entries)
    report = {
        "production_registry": str(PRODUCTION.relative_to(ROOT)),
        "candidate_registry": str(CANDIDATE_REGISTRY.relative_to(ROOT)),
        "validated_registry": str(VALIDATED_REGISTRY.relative_to(ROOT)),
        "limits": {"batch_size": MAX_BATCH_SIZE, "pages_per_company": MAX_PAGES, "normalized_jobs_per_company": MAX_JOBS},
        "totals": {
            "candidates_investigated": len(results),
            "avature_candidates": sum(entry["platform"] == "avature" for entry in CANDIDATE_ENTRIES),
            "icims_candidates": sum(entry["platform"] == "icims" for entry in CANDIDATE_ENTRIES),
            "validated_with_jobs": status_counts[validator.VALIDATED_WITH_JOBS],
            "validated_zero_jobs": status_counts[validator.VALIDATED_ZERO_JOBS],
            "excluded": sum(not result["safe_to_activate"] for result in results),
            "temporary_failures": status_counts[validator.TEMPORARY_FAILURE],
            "candidate_registry_count": candidate_count,
            "validated_registry_count": validated_count,
            "projected_original_company_coverage": 87 + len(safe_entries),
        },
        "results": results,
    }
    _write_inventory(results)
    validator._write_json_if_changed(REPORT, report)
    return report


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--timeout", type=int, default=validator.DEFAULT_TIMEOUT)
    parser.add_argument("--batch-delay", type=float, default=validator.DEFAULT_BATCH_DELAY)
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        build_candidate_registry()
        if args.dry_run:
            _write_inventory([])
            report = {
                "dry_run": True,
                "planned_companies": EXPECTED_COMPANIES,
                "batch_size": MAX_BATCH_SIZE,
                "maximum_pages": MAX_PAGES,
                "maximum_jobs": MAX_JOBS,
            }
        else:
            report = validate(timeout=args.timeout, batch_delay=args.batch_delay)
    except (OSError, ValueError, yaml.YAMLError) as exc:
        print(f"Validation setup failed: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
